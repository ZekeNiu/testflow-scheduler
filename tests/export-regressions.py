"""Generate real XLSX files and independently check elapsed-time formulas.

Requires Node.js and openpyxl. Use --output-dir to retain workbooks for Excel QA.
"""
import argparse
import ast
import json
import math
import os
from pathlib import Path
import re
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.utils.cell import range_boundaries


ROOT = Path(__file__).resolve().parents[1]
NS = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
NODE_EXPORT = r"""
const fs=require('fs'),path=require('path');
const E=require('./src/testflow-engine.js');
require('./src/testflow-xlsx-template.js');
const X=require('./src/testflow-export.js');
const fixtures=JSON.parse(fs.readFileSync(0,'utf8'));
const wall=(c,s)=>{const d=new Date(c.startMs+Math.round(s*1000));return [d.getFullYear(),d.getMonth()+1,d.getDate(),d.getHours(),d.getMinutes(),d.getSeconds(),d.getMilliseconds()];};
(async()=>{
  const results=[];
  for(const f of fixtures){
    const calculation=E.calculate(f.plan);
    if(!calculation.result)throw new Error(JSON.stringify(calculation.errors));
    const {config:c,result:r}=calculation;
    const file=path.join(process.argv[1],f.name+'.xlsx');
    fs.writeFileSync(file,await X.build(calculation));
    const reports=r.arrivalPlan.map(b=>b.report),earliest=Math.min(...reports);
    results.push({name:f.name,file,start:c.startMs,offsets:c.startMs===null?[]:[new Date(c.startMs).getTimezoneOffset(),new Date(c.startMs+r.actual*1000).getTimezoneOffset()],prep:c.prep,actual:r.actual,earliest,coverage:r.actual-Math.min(0,earliest),waves:new Set(reports.map(s=>Math.round(s*1000))).size,arrivals:r.arrivalPlan,records:r.records,stats:r.stationStats,clock:c.startMs===null?null:wall(c,c.prep+r.records[0].begin),endClock:c.startMs===null?null:wall(c,c.prep+r.records[0].end),earliestClock:c.startMs===null?null:wall(c,earliest),closeClock:c.startMs===null?null:wall(c,r.actual)});
  }
  console.log(JSON.stringify(results));
})().catch(e=>{console.error(e);process.exitCode=1;});
"""


def station(name="作业 A", **extra):
    return dict(name=name, enabled=True, kind="independent", cap=1,
                duration=3600, reset=300, gap=0, policy="full", **extra)


def plan(**extra):
    data = dict(version=4, timeUnit="sec", n=2, mode="个人流水线", groups="",
                start="2026-10-03T09:00:00", prep=60, close=120, buffer=60,
                breaks=[], arrivalMode="auto", arrivalBatchSize="", arrivalLead=300,
                stations=[station()])
    data.update(extra)
    return data


def fixture(name, **extra):
    return dict(name=name, plan=plan(**extra))


def near(actual, expected, label):
    assert math.isclose(actual, expected, rel_tol=1e-10, abs_tol=1e-9), (label, actual, expected)


def cells_xml(archive, sheet):
    tree = ET.fromstring(archive.read(f"xl/worksheets/sheet{sheet}.xml"))
    return {cell.attrib["r"]: cell for cell in tree.findall(".//x:sheetData/x:row/x:c", NS)}


def number(cells, address):
    return float(cells[address].find("x:v", NS).text)


def recalculate(cells, address):
    """Evaluate the workbook's arithmetic, using XML values rather than date objects."""
    formula = cells[address].find("x:f", NS)
    if formula is None:
        return number(cells, address)
    expression = re.sub(r"\b([A-Z]+[1-9][0-9]*)\b",
                        lambda match: repr(recalculate(cells, match[1])), formula.text)

    def evaluate(node):
        if isinstance(node, ast.Expression):
            return evaluate(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.USub):
            return -evaluate(node.operand)
        if isinstance(node, ast.BinOp):
            left, right = evaluate(node.left), evaluate(node.right)
            operations = {ast.Add: lambda: left+right, ast.Sub: lambda: left-right,
                          ast.Mult: lambda: left*right, ast.Div: lambda: left/right}
            if type(node.op) in operations:
                return operations[type(node.op)]()
        raise AssertionError(("Unsupported exported formula", formula.text))

    return evaluate(ast.parse(expression, mode="eval"))


def wall_parts(value):
    return [value.year, value.month, value.day, value.hour, value.minute,
            value.second, value.microsecond // 1000]


def find_label(sheet, label):
    for row in sheet:
        if row[0].value == label:
            return row[0].row
    raise AssertionError((sheet.title, "missing label", label))


def check_workbook(result):
    with zipfile.ZipFile(result["file"]) as archive:
        assert archive.testzip() is None
        xml_cells = {i: cells_xml(archive, i) for i in range(1, 5)}
    workbook = openpyxl.load_workbook(result["file"], data_only=False)
    cached = openpyxl.load_workbook(result["file"], data_only=True)
    assert workbook.sheetnames == ["排程总览", "人员排程", "站点排程", "到场安排"]
    overview = workbook["排程总览"]
    heading = find_label(overview, "测试站点")
    for index, stat in enumerate(result["stats"], heading+1):
        # Recalculation must preserve elapsed time even when local clocks jump.
        near(recalculate(xml_cells[1], f"F{index}") * 86400,
             stat["workDuration"], result["name"]+" station work duration")
        near(recalculate(xml_cells[1], f"G{index}"), stat["utilization"], "utilization formula")
        near(number(xml_cells[1], f"F{index}") * 86400, stat["workDuration"], "cached duration")
    for index in (2, 3):
        for row in range(6, 6+len(result["records"])):
            near(recalculate(xml_cells[index], f"J{row}") * 86400,
                 {"fractional": 30.125, "rotation": 10, "batch-tail": 10}.get(result["name"], 3600),
                 "complete test duration")
            near(recalculate(xml_cells[index], f"J{row}"),
                 number(xml_cells[index], f"J{row}"), "duration formula cache")
    assert overview["A6"].value == "开场至收尾时长"
    earliest_row = find_label(overview, "最早到场")
    close_row = find_label(overview, "收尾结束")
    coverage_row = find_label(overview, "场地覆盖时长")
    near(number(xml_cells[1], f"D{coverage_row}") * 86400, result["coverage"], "venue coverage")
    for index, sheet in enumerate(workbook, 1):
        table = list(sheet.tables.values())[0]
        first_col, first_row, last_col, last_row = range_boundaries(table.ref)
        assert table.autoFilter.ref == table.ref
        assert [column.name for column in table.tableColumns] == [
            sheet.cell(first_row, col).value for col in range(first_col, last_col+1)]
        assert all(sheet.cell(first_row, col).style_id > 0 for col in range(first_col, last_col+1))
        if index in (2, 3):
            assert sheet["H5"].value == "计划开测时间"
            assert last_row == 5+len(result["records"])
        if index > 1:
            assert sheet.freeze_panes is not None
    arrivals = workbook["到场安排"]
    assert arrivals["A5"].value == "到场波次"
    assert f'{result["waves"]} 个到场波次' in arrivals["A3"].value
    assert f'{len(result["arrivals"])} 个接待组' in arrivals["A3"].value
    reports = sorted(set(round(b["report"] * 1000) for b in result["arrivals"]))
    for row, arrival in enumerate(sorted(result["arrivals"], key=lambda b: (b["report"], b["batch"])), 6):
        assert arrivals.cell(row, 1).value == reports.index(round(arrival["report"]*1000))+1
        assert f'接待组 {arrival["batch"]}' in arrivals.cell(row, 2).value
        assert arrivals.cell(row, 3).value == len(arrival["people"])
        assert arrivals.cell(row, 4).value == (arrival["group"] or "—")
    assert arrivals.max_row == 5+len(result["arrivals"])
    if result["name"] == "batch-tail":
        assert any(cell.value == "满批优先，尾批可不足" for row in overview for cell in row)
        assert [record["unit"] for record in result["records"]] == [1, 1, 2]
    absolute = result["start"] is not None and result["earliestClock"][0] >= 1900
    if absolute:
        assert wall_parts(cached["人员排程"]["H6"].value) == result["clock"]
        assert wall_parts(cached["人员排程"]["I6"].value) == result["endClock"]
        assert wall_parts(cached["排程总览"].cell(earliest_row, 4).value) == result["earliestClock"]
        assert wall_parts(cached["排程总览"].cell(close_row, 4).value) == result["closeClock"]
    else:
        near(number(xml_cells[1], f"D{earliest_row}"), result["earliest"], "elapsed earliest arrival")
        near(number(xml_cells[2], "H6"), result["prep"]+result["records"][0]["begin"], "elapsed time")
    workbook.close()
    cached.close()


def run(output):
    groups = {
        "America/New_York": [fixture("spring-dst", start="2026-03-08T01:30:00"),
                             fixture("fall-dst", start="2026-11-01T01:30:00")],
        "Asia/Shanghai": [fixture("shanghai"),
                          fixture("late-first-arrival", arrivalLead=0),
                          fixture("excel-1900", n=1, start="1900-02-28T23:30:00", prep=0,
                                  arrivalMode="all", arrivalLead=0),
                          fixture("before-1900", n=1, start="1899-12-31T23:30:00", prep=0,
                                  arrivalMode="all", arrivalLead=0),
                          fixture("elapsed", start=""),
                          fixture("fractional", n=3, prep=15.5, close=2.25, buffer=1.25,
                                  arrivalLead=30.75,
                                  stations=[dict(station(), duration=30.125, reset=0.875)]),
                          fixture("rotation", n=6, mode="分组轮转", groups=3, arrivalMode="all",
                                  stations=[dict(station(name), cap=2, duration=10, reset=0)
                                            for name in ("A", "B", "C")]),
                          fixture("batch-tail", n=3, stations=[dict(station(), kind="batch",
                                  cap=2, duration=10, reset=2)])],
    }
    total = 0
    for timezone, fixtures in groups.items():
        env = dict(os.environ, TZ=timezone)
        generated = subprocess.run(["node", "-e", NODE_EXPORT, str(output)], cwd=ROOT,
                                   env=env, input=json.dumps(fixtures, ensure_ascii=False),
                                   text=True, encoding="utf-8", capture_output=True, check=True)
        for result in json.loads(generated.stdout):
            if result["name"] in ("spring-dst", "fall-dst"):
                assert result["offsets"][0] != result["offsets"][1], "DST fixture did not cross the clock change"
            check_workbook(result)
            total += 1
    print(f"XLSX regression checks passed: {total} workbooks, DST/Shanghai/1900/elapsed/fractional/waves, independent formula recomputation.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    if args.output_dir:
        args.output_dir.mkdir(parents=True, exist_ok=True)
        run(args.output_dir.resolve())
    else:
        with tempfile.TemporaryDirectory(prefix="testflow-export-") as folder:
            run(Path(folder))
